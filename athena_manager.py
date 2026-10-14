import copy
import csv
from datetime import datetime, timezone
import io
import json
import logging
import math
import os
import re
import socket
import tempfile
import threading
import time
import uuid
from typing import Any, Dict, List, Optional, Set, Tuple
import boto3
from boto3.s3.transfer import TransferConfig
import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import sqlglot
from sqlglot import exp
import xlrd
from botocore.exceptions import ClientError

from database import (
    AthenaExecutionError,
    AthenaResultRetrievalError,
    AthenaTimeoutError,
    SQLValidationError,
    validate_sql_query,
)

logger = logging.getLogger(__name__)

# Authoritative schema overrides for standard AWS Cost & Usage Report (CUR) columns
CUR_COLUMN_TYPE_OVERRIDES: Dict[str, str] = {
    "line_item_unblended_cost": "DOUBLE",
    "line_item_blended_cost": "DOUBLE",
    "line_item_unblended_rate": "DOUBLE",
    "line_item_blended_rate": "DOUBLE",
    "line_item_usage_amount": "DOUBLE",
    "pricing_public_on_demand_cost": "DOUBLE",
    "pricing_public_on_demand_rate": "DOUBLE",
    "reservation_effective_cost": "DOUBLE",
    "savings_plan_savings_plan_effective_cost": "DOUBLE",
    "line_item_usage_start_date": "STRING",
    "line_item_usage_end_date": "STRING",
    "bill_billing_period_start_date": "STRING",
    "bill_billing_period_end_date": "STRING",
    "line_item_product_code": "STRING",
    "line_item_resource_id": "STRING",
    "line_item_usage_account_id": "STRING",
    "line_item_line_item_type": "STRING",
    "line_item_operation": "STRING",
    "line_item_usage_type": "STRING",
    "product_instance_type": "STRING",
    "product_product_name": "STRING",
    "product_region": "STRING",
}

ALLOWED_ATHENA_PRIMITIVE_TYPES: Set[str] = {
    "STRING", "BIGINT", "INT", "DOUBLE", "FLOAT", "BOOLEAN", "DATE", "TIMESTAMP"
}

NUMERIC_ATHENA_TYPES: Set[str] = {"BIGINT", "INT", "DOUBLE", "FLOAT"}

COMPATIBLE_PARQUET_TYPE_OVERRIDES: Dict[str, Set[str]] = {
    "INT": {"INT", "BIGINT"},
    "BIGINT": {"BIGINT"},
    "FLOAT": {"FLOAT", "DOUBLE"},
    "DOUBLE": {"DOUBLE"},
    "BOOLEAN": {"BOOLEAN"},
    "DATE": {"DATE"},
    "TIMESTAMP": {"TIMESTAMP"},
    "STRING": {"STRING"},
}

GLUE_READ_ONLY_TABLE_FIELDS: Set[str] = {
    "DatabaseName",
    "CreateTime",
    "UpdateTime",
    "CreatedBy",
    "IsRegisteredWithLakeFormation",
    "TargetTable",
    "CatalogId",
    "VersionId",
    "FederatedTable",
    "IsMultiDialectView",
}

MAX_EXCEL_UPLOAD_BYTES = 50 * 1024 * 1024  # 50 MB compressed upload cap
MAX_EXCEL_ROWS = 100_000  # Decompressed row cap to bound container RAM
MAX_EXCEL_COLS = 200      # Decompressed column cap to bound container RAM

ACTIVE_INGESTION_STATES: Set[str] = {"UPLOADING", "VALIDATING", "CUTOVER", "ACTIVE"}
TERMINAL_INGESTION_STATES: Set[str] = {"RETIRED", "FAILED"}
STAGING_VALIDATION_ROW_LIMIT = 1000


def _validate_athena_decimal_type(type_str: str) -> bool:
    """Validates Athena DECIMAL(precision, scale) constraints: 1 <= precision <= 38 and 0 <= scale <= precision."""
    match = re.fullmatch(r"DECIMAL\((\d{1,2}),\s*(\d{1,2})\)", type_str.strip().upper())
    if not match:
        return False
    precision, scale = int(match.group(1)), int(match.group(2))
    return 1 <= precision <= 38 and 0 <= scale <= precision


def _is_valid_athena_ddl_type(type_str: str) -> bool:
    upper_t = type_str.strip().upper()
    return upper_t in ALLOWED_ATHENA_PRIMITIVE_TYPES or _validate_athena_decimal_type(upper_t)


def _is_numeric_athena_type(type_str: str) -> bool:
    upper_t = type_str.strip().upper()
    return upper_t in NUMERIC_ATHENA_TYPES or _validate_athena_decimal_type(upper_t)


def configure_workgroup_for_deployment(
    work_group: str = "nl-bi-finops-workgroup",
    max_scan_mb: int = 100,
    region: str = "us-east-1",
) -> None:
    """
    One-time deployment/provisioning helper to configure the Athena Workgroup
    scan cutoff. Validates arguments before calling AWS. Not invoked at application
    startup so the runtime IAM role only requires athena:GetWorkGroup.
    """
    if not work_group or not re.fullmatch(r"[a-zA-Z0-9._-]{1,128}", work_group):
        raise ValueError(f"Invalid Athena Workgroup name: {work_group!r}")
    if not isinstance(max_scan_mb, int) or isinstance(max_scan_mb, bool) or max_scan_mb <= 0:
        raise ValueError(f"max_scan_mb must be a positive integer, got {max_scan_mb!r}")

    athena = boto3.client("athena", region_name=region)
    athena.update_work_group(
        WorkGroup=work_group,
        ConfigurationUpdates={
            "EnforceWorkGroupConfiguration": True,
            "BytesScannedCutoffPerQuery": max_scan_mb * 1024 * 1024,
        },
    )


class AthenaManager:

    # Local thread lock serves only as an in-process first gate before distributed DynamoDB/S3 coordination
    _REGISTRY_LOCK = threading.Lock()
    _TABLE_LOCKS: Dict[str, threading.Lock] = {}
    _INGESTION_STATES: Dict[str, Dict[str, Any]] = {}

    def __init__(
        self,
        database: str,
        bucket_name: str,
        region: str = "us-east-1",
        work_group: Optional[str] = None,
        max_scan_mb: int = 100,
        query_timeout_sec: int = 30,
        strict_workgroup_check: bool = True,
        lock_table_name: Optional[str] = None,
        lock_lease_seconds: int = 300,
        lease_heartbeat_interval_sec: Optional[float] = None,
        state_retention_seconds: int = 7 * 86400,
        task_owner_id: Optional[str] = None,
        use_atomic_glue_cutover: bool = False,
        validate_full_dataset: bool = False,
    ):
        if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_]*", database):
            raise ValueError(f"Invalid Glue database identifier: {database}")
        if not isinstance(max_scan_mb, int) or isinstance(max_scan_mb, bool) or max_scan_mb <= 0:
            raise ValueError(f"max_scan_mb must be a positive integer, got {max_scan_mb!r}")
        if not isinstance(query_timeout_sec, (int, float)) or isinstance(query_timeout_sec, bool) or query_timeout_sec <= 0:
            raise ValueError(f"query_timeout_sec must be a positive number, got {query_timeout_sec!r}")
        if not isinstance(lock_lease_seconds, int) or lock_lease_seconds < 30:
            raise ValueError("lock_lease_seconds must be an integer >= 30.")

        resolved_wg = work_group or os.getenv("ATHENA_WORKGROUP", "nl-bi-finops-workgroup")
        if not re.fullmatch(r"[a-zA-Z0-9._-]{1,128}", resolved_wg):
            raise ValueError(f"Invalid Athena Workgroup name: {resolved_wg!r}")

        self.database = database
        self.bucket_name = bucket_name
        self.region = region
        self.work_group = resolved_wg
        self.max_scan_mb = max_scan_mb
        self.max_scan_bytes = max_scan_mb * 1024 * 1024
        self.query_timeout_sec = query_timeout_sec
        self.strict_workgroup_check = strict_workgroup_check
        self.lock_table_name = lock_table_name or os.getenv(
            "INGESTION_LOCK_TABLE", "nl_bi_finops_ingestion_locks"
        )
        self.lock_lease_seconds = lock_lease_seconds
        self.lease_heartbeat_interval_sec = (
            float(lease_heartbeat_interval_sec)
            if lease_heartbeat_interval_sec is not None
            else float(max(5, lock_lease_seconds // 3))
        )
        self.state_retention_seconds = state_retention_seconds
        # Every manager instance gets a unique worker identity so two workers inside
        # the same ECS task (sharing ECS_TASK_ID) cannot share or overwrite a lease.
        task_identity = task_owner_id or os.getenv("ECS_TASK_ID") or socket.gethostname()
        self.task_owner_id = f"{task_identity}:{os.getpid()}:{uuid.uuid4().hex}"
        self.use_atomic_glue_cutover = use_atomic_glue_cutover
        self.validate_full_dataset = validate_full_dataset
        self.output_location = f"s3://{bucket_name}/athena-results/"
        self.workgroup_guardrail_verified = False

        self.athena = boto3.client("athena", region_name=region)
        self.s3 = boto3.client("s3", region_name=region)
        self.glue = boto3.client("glue", region_name=region)
        try:
            self.dynamodb = boto3.client("dynamodb", region_name=region)
        except Exception:
            self.dynamodb = None
        self.dialect = "trino"

        self.transfer_config = TransferConfig(
            multipart_threshold=8 * 1024 * 1024,
            multipart_chunksize=8 * 1024 * 1024,
            max_concurrency=4,
            use_threads=True,
        )

        self._verify_workgroup_guardrails(strict=self.strict_workgroup_check)

    @classmethod
    def _get_table_lock(cls, table_key: str) -> threading.Lock:
        with cls._REGISTRY_LOCK:
            if table_key not in cls._TABLE_LOCKS:
                cls._TABLE_LOCKS[table_key] = threading.Lock()
            return cls._TABLE_LOCKS[table_key]

    # -------------------------------------------------------------------------
    # Distributed Concurrency Locking & Durable State Persistence (P0 / P1)
    # -------------------------------------------------------------------------

    def _state_manifest_s3_key(self, table_name: str, ingestion_id: str) -> str:
        return f"_ingestion_state/{self.database}/{table_name}/{ingestion_id}.json"

    def _acquire_distributed_table_lock(self, table_name: str, ingestion_id: str) -> None:
        """Acquire a fail-closed DynamoDB lease for this exact worker + ingestion."""
        if self.dynamodb is None or not self.lock_table_name:
            raise AthenaExecutionError(
                "Fail-Closed Distributed Lock Error: DynamoDB is unavailable or lock table is unset."
            )
        lock_key = f"{self.database}.{table_name}"
        now = int(time.time())
        expires_at = now + self.lock_lease_seconds
        try:
            self.dynamodb.put_item(
                TableName=self.lock_table_name,
                Item={
                    "lock_key": {"S": lock_key},
                    "owner_id": {"S": self.task_owner_id},
                    "ingestion_id": {"S": ingestion_id},
                    "state": {"S": "ACQUIRED"},
                    "lease_expires_at": {"N": str(expires_at)},
                    "updated_at": {"N": str(now)},
                },
                ConditionExpression="attribute_not_exists(lock_key) OR lease_expires_at < :now",
                ExpressionAttributeValues={":now": {"N": str(now)}},
            )
        except ClientError as e:
            code = e.response.get("Error", {}).get("Code", "")
            if code == "ConditionalCheckFailedException":
                raise AthenaExecutionError(
                    f"Concurrent distributed ingestion conflict for '{lock_key}'."
                ) from e
            raise AthenaExecutionError(
                f"Fail-Closed Distributed Lock Error: Cannot acquire required DynamoDB lock for '{lock_key}' ({code}): {e}"
            ) from e

    def _renew_distributed_table_lock(self, table_name: str, ingestion_id: str, state: str) -> None:
        """Renew only if both worker and ingestion still own the unexpired lease."""
        if self.dynamodb is None or not self.lock_table_name:
            raise AthenaExecutionError("Cannot renew distributed lock: DynamoDB is unavailable.")
        lock_key = f"{self.database}.{table_name}"
        now = int(time.time())
        try:
            self.dynamodb.update_item(
                TableName=self.lock_table_name,
                Key={"lock_key": {"S": lock_key}},
                UpdateExpression="SET lease_expires_at=:exp, updated_at=:now, #st=:state",
                ConditionExpression="owner_id=:owner AND ingestion_id=:ing AND lease_expires_at >= :now",
                ExpressionAttributeNames={"#st": "state"},
                ExpressionAttributeValues={
                    ":exp": {"N": str(now + self.lock_lease_seconds)},
                    ":now": {"N": str(now)},
                    ":state": {"S": state},
                    ":owner": {"S": self.task_owner_id},
                    ":ing": {"S": ingestion_id},
                },
            )
        except ClientError as e:
            raise AthenaExecutionError(
                f"Distributed lease lost or expired for '{lock_key}' / ingestion '{ingestion_id}' "
                f"during '{state}'; aborting to prevent split-brain: {e}"
            ) from e

    def _release_distributed_table_lock(self, table_name: str, ingestion_id: str) -> None:
        """Release only this ingestion's lease; never delete another worker's lock."""
        if self.dynamodb is None or not self.lock_table_name:
            return
        lock_key = f"{self.database}.{table_name}"
        try:
            self.dynamodb.delete_item(
                TableName=self.lock_table_name,
                Key={"lock_key": {"S": lock_key}},
                ConditionExpression="owner_id=:owner AND ingestion_id=:ing",
                ExpressionAttributeValues={
                    ":owner": {"S": self.task_owner_id},
                    ":ing": {"S": ingestion_id},
                },
            )
        except ClientError as e:
            if e.response.get("Error", {}).get("Code") != "ConditionalCheckFailedException":
                logger.error("Failed to release DynamoDB lock %s: %s", lock_key, e)

    def _set_ingestion_state(
        self, s3_prefix: str, table_name: str, ingestion_id: str, state: str
    ) -> Dict[str, Any]:
        """Persist authoritative state to S3; fail closed if persistence fails."""
        now = time.time()
        with self._REGISTRY_LOCK:
            existing = self._INGESTION_STATES.get(s3_prefix)
            record = {
                "ingestion_id": ingestion_id,
                "database": self.database,
                "table_name": table_name,
                "s3_prefix": s3_prefix,
                "state": state,
                "owner_id": self.task_owner_id,
                "started_at": existing["started_at"] if existing else now,
                "updated_at": now,
                "lease_expires_at": now + self.lock_lease_seconds,
                "ttl_expires_at": now + self.state_retention_seconds,
            }
        key = self._state_manifest_s3_key(table_name, ingestion_id)
        try:
            self.s3.put_object(
                Bucket=self.bucket_name,
                Key=key,
                Body=json.dumps(record).encode("utf-8"),
                ContentType="application/json",
            )
        except Exception as e:
            raise AthenaExecutionError(
                f"Fail-Closed State Persistence Error: Cannot persist required ingestion state manifest '{key}'; aborting: {e}"
            ) from e
        with self._REGISTRY_LOCK:
            self._INGESTION_STATES[s3_prefix] = record
        if state in ("UPLOADING", "VALIDATING", "CUTOVER", "ACTIVE"):
            self._renew_distributed_table_lock(table_name, ingestion_id, state)
        return record

    def _load_durable_ingestion_states(self, table_name: str) -> Dict[str, Dict[str, Any]]:
        """Load all manifests or fail closed; incomplete state cannot authorize cleanup."""
        found: Dict[str, Dict[str, Any]] = {}
        prefix = f"_ingestion_state/{self.database}/{table_name}/"
        try:
            paginator = self.s3.get_paginator("list_objects_v2")
            for page in paginator.paginate(Bucket=self.bucket_name, Prefix=prefix):
                if not isinstance(page, dict):
                    continue
                for obj in page.get("Contents", []):
                    key = obj.get("Key", "")
                    if not key.endswith(".json"):
                        continue
                    resp = self.s3.get_object(Bucket=self.bucket_name, Key=key)
                    data = json.loads(resp["Body"].read().decode("utf-8"))
                    if not isinstance(data, dict) or not data.get("s3_prefix") or not data.get("state"):
                        raise ValueError(f"Malformed ingestion manifest: {key}")
                    found[data["s3_prefix"]] = data
        except Exception as e:
            raise AthenaExecutionError(
                f"Cannot establish authoritative ingestion state under s3://{self.bucket_name}/{prefix}: {e}"
            ) from e
        with self._REGISTRY_LOCK:
            for k, v in self._INGESTION_STATES.items():
                if v.get("table_name") == table_name:
                    old = found.get(k)
                    if old is None or float(v.get("updated_at", 0)) > float(old.get("updated_at", 0)):
                        found[k] = dict(v)
        return found

    def get_ingestion_state(
        self, s3_prefix: str, table_name: Optional[str] = None
    ) -> Optional[Dict[str, Any]]:
        with self._REGISTRY_LOCK:
            info = self._INGESTION_STATES.get(s3_prefix)
            if info:
                return dict(info)

        inferred_table = table_name
        if not inferred_table:
            m = re.match(r"^data/([^/]+)/ingestion_id=([a-f0-9]+)/", s3_prefix)
            if m:
                inferred_table = m.group(1)

        if inferred_table:
            durable = self._load_durable_ingestion_states(inferred_table)
            if s3_prefix in durable:
                return dict(durable[s3_prefix])
        return None

    def prune_stale_ingestion_records(
        self, table_name: str, retention_seconds: Optional[int] = None
    ) -> List[str]:
        """
        Enforces a lifecycle retention policy for terminal ingestion state records
        (RETIRED, FAILED, or superseded ACTIVE records) so state does not accumulate
        indefinitely in memory or S3.
        """
        max_age = retention_seconds if retention_seconds is not None else self.state_retention_seconds
        now_ts = time.time()
        snapshot = self._get_existing_table_snapshot(table_name)
        active_prefix = snapshot.get("s3_prefix") if snapshot else None

        durable_states = self._load_durable_ingestion_states(table_name)
        pruned_prefixes: List[str] = []

        for prefix, rec in list(durable_states.items()):
            state = rec.get("state", "")
            updated_at = float(rec.get("updated_at", now_ts))
            age = now_ts - updated_at
            is_terminal_or_superseded = (
                state in TERMINAL_INGESTION_STATES
                or (state == "ACTIVE" and prefix != active_prefix)
            )
            if is_terminal_or_superseded and age >= max_age:
                ing_id = rec.get("ingestion_id", "")
                if ing_id:
                    manifest_key = self._state_manifest_s3_key(table_name, ing_id)
                    try:
                        self.s3.delete_objects(
                            Bucket=self.bucket_name,
                            Delete={"Objects": [{"Key": manifest_key}], "Quiet": False},
                        )
                    except Exception as e:
                        logger.warning(f"Could not delete expired state manifest '{manifest_key}': {e}")

                with self._REGISTRY_LOCK:
                    self._INGESTION_STATES.pop(prefix, None)
                pruned_prefixes.append(prefix)

        return pruned_prefixes

    def _verify_workgroup_guardrails(self, strict: bool = True) -> bool:
        """
        Verifies via get_work_group() that the target Athena Workgroup has
        EnforceWorkGroupConfiguration=True and BytesScannedCutoffPerQuery <= self.max_scan_bytes.
        """
        try:
            wg_resp = self.athena.get_work_group(WorkGroup=self.work_group)
            wg_config = wg_resp.get("WorkGroup", {}).get("Configuration", {})
            actual_cutoff = wg_config.get("BytesScannedCutoffPerQuery")
            enforced = wg_config.get("EnforceWorkGroupConfiguration", False)

            self.workgroup_guardrail_verified = (
                isinstance(actual_cutoff, int)
                and 0 < actual_cutoff <= self.max_scan_bytes
                and enforced is True
            )
            if not self.workgroup_guardrail_verified:
                msg = (
                    f"Fail-Closed FinOps Guardrail Error: Athena Workgroup '{self.work_group}' "
                    f"does not have an enforced scan cutoff <= {self.max_scan_bytes} bytes. "
                    f"Got BytesScannedCutoffPerQuery={actual_cutoff}, EnforceWorkGroupConfiguration={enforced}."
                )
                logger.error(msg)
                if strict:
                    raise AthenaExecutionError(msg)
            return self.workgroup_guardrail_verified
        except ClientError as e:
            msg = (
                f"Fail-Closed FinOps Guardrail Error: Unable to verify Athena Workgroup "
                f"'{self.work_group}' scan cutoff configuration via GetWorkGroup: {e}"
            )
            logger.error(msg)
            if strict:
                raise AthenaExecutionError(msg) from e
            return False

    @staticmethod
    def _is_retryable_athena_sql_error(status_dict: Dict[str, Any], reason: str) -> bool:
        non_retryable_markers = (
            "ACCESS_DENIED",
            "AccessDenied",
            "InsufficientLakeFormationPermissions",
            "Query exhausted",
            "Scanned data limit",
            "BytesScannedCutoffPerQuery",
            "data scanned limit",
            "scan limit",
            "scan cutoff",
            "WorkGroup",
            "INTERNAL_ERROR",
            "THROTTLING",
            "ThrottlingException",
            "SlowDown",
            "HIVE_METASTORE_ERROR",
            "HIVE_CANNOT_OPEN_SPLIT",
            "HIVE_BAD_DATA",
            "HIVE_CURSOR_ERROR",
            "S3",
        )
        if any(m.lower() in reason.lower() for m in non_retryable_markers):
            return False

        athena_err = status_dict.get("AthenaError")
        if isinstance(athena_err, dict) and athena_err:
            error_category = athena_err.get("ErrorCategory")
            error_type = athena_err.get("ErrorType")
            err_msg = str(athena_err.get("ErrorMessage") or reason)

            if any(m.lower() in err_msg.lower() for m in non_retryable_markers):
                return False

            if error_category in (1, 3):
                return False

            non_sql_user_error_types = {1003, 1004, 1006, 1007}
            if error_type in non_sql_user_error_types:
                return False

            if error_category == 2:
                return True

        conservative_sql_codes = (
            "SYNTAX_ERROR:",
            "COLUMN_NOT_FOUND:",
            "TABLE_NOT_FOUND:",
            "TYPE_MISMATCH:",
            "FUNCTION_NOT_FOUND:",
            "INVALID_CAST_ARGUMENT:",
            "AMBIGUOUS_NAME:",
            "EXPRESSION_NOT_AGGREGATE:",
        )
        return any(code in reason for code in conservative_sql_codes)

    def _poll_athena_execution(self, query_execution_id: str) -> dict:
        """Polls an Athena QueryExecutionId using monotonic time with a hard timeout circuit breaker."""
        start_time = time.monotonic()
        while True:
            if (time.monotonic() - start_time) > self.query_timeout_sec:
                try:
                    self.athena.stop_query_execution(QueryExecutionId=query_execution_id)
                except Exception:
                    pass
                raise AthenaTimeoutError(
                    f"Athena query (Execution ID: {query_execution_id}) exceeded timeout limit "
                    f"of {self.query_timeout_sec}s and was cancelled."
                )

            status_resp = self.athena.get_query_execution(
                QueryExecutionId=query_execution_id
            )
            query_info = status_resp.get("QueryExecution", {})
            status_dict = query_info.get("Status", {})
            state = status_dict.get("State")

            if state == "SUCCEEDED":
                return query_info
            elif state in ["FAILED", "CANCELLED"]:
                reason = status_dict.get("StateChangeReason", "Unknown error")
                if state == "FAILED" and self._is_retryable_athena_sql_error(status_dict, reason):
                    raise SQLValidationError(
                        f"Athena SQL Error (Execution ID: {query_execution_id}): {reason}"
                    )
                raise AthenaExecutionError(
                    f"Athena Query {state} (Execution ID: {query_execution_id}): {reason}"
                )
            time.sleep(0.5)

    def _has_active_cutover_in_progress(self) -> bool:
        """Checks whether any table currently has an active CUTOVER state in progress."""
        now_ts = time.time()
        with self._REGISTRY_LOCK:
            for rec in self._INGESTION_STATES.values():
                if (
                    rec.get("state") == "CUTOVER"
                    and float(rec.get("lease_expires_at", 0)) > now_ts
                ):
                    return True
        return False

    def execute_query(
        self, sql: str, retry_on_cutover: bool = True
    ) -> Tuple[pd.DataFrame, str]:
        """
        Strictly read-only analytical query execution boundary for Amazon Athena.
        """
        if self.strict_workgroup_check and not self.workgroup_guardrail_verified:
            raise AthenaExecutionError(
                f"Fail-Closed FinOps Guardrail: Refusing to execute Athena query because "
                f"Workgroup '{self.work_group}' {self.max_scan_mb} MB scan cutoff is not verified."
            )

        validate_sql_query(sql, dialect="trino")
        sql_clean = sql.strip()
        if sql_clean.endswith(";"):
            sql_clean = sql_clean[:-1].strip()

        max_attempts = 2 if retry_on_cutover else 1
        query_execution_id = ""
        query_info: Dict[str, Any] = {}

        for attempt in range(1, max_attempts + 1):
            try:
                response = self.athena.start_query_execution(
                    QueryString=sql_clean,
                    QueryExecutionContext={"Database": self.database},
                    ResultConfiguration={"OutputLocation": self.output_location},
                    WorkGroup=self.work_group,
                )
            except ClientError as e:
                raise AthenaExecutionError(f"Failed to start Athena query execution: {e}") from e

            query_execution_id = response["QueryExecutionId"]
            try:
                query_info = self._poll_athena_execution(query_execution_id)
                break
            except (SQLValidationError, AthenaExecutionError) as err:
                is_table_missing = "TABLE_NOT_FOUND" in str(err) or "does not exist" in str(err)
                if attempt < max_attempts and is_table_missing and self._has_active_cutover_in_progress():
                    logger.info(
                        "Query encountered transient TABLE_NOT_FOUND during active table cutover; "
                        "waiting briefly and retrying once."
                    )
                    time.sleep(0.2)
                    continue
                raise

        output_location = query_info.get("ResultConfiguration", {}).get(
            "OutputLocation", ""
        )
        if not output_location.startswith("s3://") or not output_location.endswith(".csv"):
            raise AthenaResultRetrievalError(
                f"Unexpected or missing Athena CSV output location '{output_location}' "
                f"for QueryExecutionId '{query_execution_id}'."
            )

        s3_path = output_location.replace("s3://", "", 1)
        if "/" not in s3_path:
            raise AthenaResultRetrievalError(
                f"Malformed S3 output URI '{output_location}' for QueryExecutionId '{query_execution_id}'."
            )
        target_bucket, result_key = s3_path.split("/", 1)

        try:
            obj = self.s3.get_object(Bucket=target_bucket, Key=result_key)
            df = pd.read_csv(obj["Body"])
        except pd.errors.EmptyDataError:
            df = pd.DataFrame()
        except ClientError as e:
            raise AthenaResultRetrievalError(
                f"Athena query '{query_execution_id}' succeeded, but reading result object "
                f"'s3://{target_bucket}/{result_key}' failed: {e}"
            ) from e
        except Exception as e:
            raise AthenaResultRetrievalError(
                f"Failed to parse result CSV for Athena QueryExecutionId '{query_execution_id}': {e}"
            ) from e

        bytes_scanned = (
            query_info.get("Statistics", {}).get("DataScannedInBytes", 0)
        )
        plan_summary = (
            f"Athena Execution ID: {query_execution_id} "
            f"(Data Scanned: {bytes_scanned / (1024 * 1024):.2f} MB)"
        )
        return df, plan_summary

    def _validate_staging_table(
        self,
        staging_table_name: str,
        col_defs: List[Tuple[str, str]],
        validate_full_dataset: Optional[bool] = None,
    ) -> None:
        """
        Validate all declared columns and numeric casts over a bounded 1,000-row sample
        (or across the full dataset when validate_full_dataset=True).
        Athena CAST and IS_NAN / IS_INFINITE checks reject malformed numeric values.
        """
        if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_]*", staging_table_name):
            raise ValueError(f"Invalid staging table identifier: {staging_table_name}")
        if not col_defs:
            raise ValueError("Cannot validate staging table with empty column definitions.")

        full_scan = (
            self.validate_full_dataset
            if validate_full_dataset is None
            else bool(validate_full_dataset)
        )

        names: List[str] = []
        checks: List[str] = ["COUNT(1) AS _sampled_rows"]
        numeric_col_indices: List[Tuple[int, str]] = []

        for idx, (name, dtype) in enumerate(col_defs):
            if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_]*", name):
                raise ValueError(f"Invalid column identifier: {name}")
            names.append(name)
            checks.append(f"COUNT({name}) AS _nn_chk_{idx}")
            if _is_numeric_athena_type(dtype):
                numeric_col_indices.append((idx, name))
                checks.append(f"SUM(CAST({name} AS DOUBLE)) AS _num_chk_{idx}")
                checks.append(
                    f"SUM(CASE WHEN IS_NAN(CAST({name} AS DOUBLE)) "
                    f"OR IS_INFINITE(CAST({name} AS DOUBLE)) THEN 1 ELSE 0 END) AS _nan_inf_chk_{idx}"
                )

        limit_clause = "" if full_scan else f" LIMIT {STAGING_VALIDATION_ROW_LIMIT}"
        inner = f"SELECT {', '.join(names)} FROM {self.database}.{staging_table_name}{limit_clause}"
        sql = f"SELECT {', '.join(checks)} FROM ({inner}) _stg_sample"
        validate_sql_query(sql, dialect="trino")

        try:
            resp = self.athena.start_query_execution(
                QueryString=sql,
                QueryExecutionContext={"Database": self.database},
                ResultConfiguration={"OutputLocation": self.output_location},
                WorkGroup=self.work_group,
            )
            info = self._poll_athena_execution(resp["QueryExecutionId"])
        except Exception as e:
            raise AthenaExecutionError(
                f"Staging table validation query failed for '{self.database}.{staging_table_name}': {e}"
            ) from e

        out = (info.get("ResultConfiguration") or {}).get("OutputLocation") if isinstance(info, dict) else None
        if isinstance(out, str) and out.startswith("s3://") and out.endswith(".csv"):
            bucket_key = out[5:].split("/", 1)
            if len(bucket_key) == 2:
                try:
                    obj = self.s3.get_object(Bucket=bucket_key[0], Key=bucket_key[1])
                    body = obj.get("Body") if isinstance(obj, dict) else None
                    if hasattr(body, "read"):
                        df = pd.read_csv(body)
                        if not df.empty:
                            if "_sampled_rows" in df.columns and int(df["_sampled_rows"].iloc[0]) <= 0:
                                raise ValueError(
                                    f"Staging table '{staging_table_name}' contains no readable rows."
                                )
                            for idx, col_name in numeric_col_indices:
                                nan_col = f"_nan_inf_chk_{idx}"
                                sum_col = f"_num_chk_{idx}"
                                if nan_col in df.columns and pd.notna(df[nan_col].iloc[0]):
                                    if int(float(df[nan_col].iloc[0])) > 0:
                                        raise ValueError(
                                            f"Numeric column '{col_name}' contains NaN or Infinite values."
                                        )
                                if sum_col in df.columns and pd.notna(df[sum_col].iloc[0]):
                                    val = float(df[sum_col].iloc[0])
                                    if not math.isfinite(val):
                                        raise ValueError(
                                            f"Numeric validation produced NaN/Infinity for column '{col_name}'."
                                        )
                except ValueError as ve:
                    raise AthenaExecutionError(
                        f"Staging table validation failed for '{self.database}.{staging_table_name}': {ve}"
                    ) from ve
                except ClientError as ce:
                    raise AthenaExecutionError(
                        f"Staging table validation S3 read failed for '{self.database}.{staging_table_name}': {ce}"
                    ) from ce

    def _atomic_glue_update_table(
        self,
        table_name: str,
        col_defs: List[Tuple[str, str]],
        storage_format: str,
        s3_location: str,
        field_delimiter: str = ",",
        existing_snapshot: Optional[Dict[str, Any]] = None,
    ) -> None:
        """Update storage location/schema while preserving existing Glue table metadata & PartitionKeys."""
        try:
            if existing_snapshot and isinstance(existing_snapshot.get("raw_table_definition"), dict):
                current = copy.deepcopy(existing_snapshot["raw_table_definition"])
            else:
                response = self.glue.get_table(DatabaseName=self.database, Name=table_name)
                current = copy.deepcopy(response.get("Table", {}))

            table_input = {
                k: v for k, v in current.items() if k not in GLUE_READ_ONLY_TABLE_FIELDS
            }
            table_input["Name"] = table_name
            table_input.setdefault("TableType", "EXTERNAL_TABLE")

            sd = dict(table_input.get("StorageDescriptor") or {})
            sd["Columns"] = [{"Name": n, "Type": t.lower()} for n, t in col_defs]
            sd["Location"] = s3_location
            if storage_format == "PARQUET":
                sd["InputFormat"] = "org.apache.hadoop.hive.ql.io.parquet.MapredParquetInputFormat"
                sd["OutputFormat"] = "org.apache.hadoop.hive.ql.io.parquet.MapredParquetOutputFormat"
                sd["SerdeInfo"] = {
                    **(sd.get("SerdeInfo") or {}),
                    "SerializationLibrary": "org.apache.hadoop.hive.ql.io.parquet.serde.ParquetHiveSerDe",
                }
            else:
                sd["InputFormat"] = "org.apache.hadoop.mapred.TextInputFormat"
                sd["OutputFormat"] = "org.apache.hadoop.hive.ql.io.HiveIgnoreKeyTextOutputFormat"
                serde = dict(sd.get("SerdeInfo") or {})
                serde["SerializationLibrary"] = "org.apache.hadoop.hive.serde2.lazy.LazySimpleSerDe"
                params = dict(serde.get("Parameters") or {})
                delim = "\t" if field_delimiter in ("\t", "\\t") else field_delimiter
                params.update({"field.delim": delim, "serialization.format": delim})
                serde["Parameters"] = params
                sd["SerdeInfo"] = serde
            table_input["StorageDescriptor"] = sd
            self.glue.update_table(DatabaseName=self.database, TableInput=table_input)
        except Exception as e:
            raise AthenaExecutionError(
                f"Atomic Glue UpdateTable failed for '{self.database}.{table_name}': {e}"
            ) from e

    def _build_and_validate_create_ddl(
        self,
        table_name: str,
        col_defs: List[Tuple[str, str]],
        storage_format: str,
        s3_location: str,
        field_delimiter: str = ",",
        expected_base_table: Optional[str] = None,
    ) -> str:
        if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_]*", self.database):
            raise ValueError(f"Invalid database identifier: {self.database}")
        if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_]*", table_name):
            raise ValueError(f"Invalid table identifier: {table_name}")
        if not col_defs:
            raise ValueError("Column definitions are required for CREATE EXTERNAL TABLE.")

        base_tbl = expected_base_table or table_name
        expected_s3_prefix = f"s3://{self.bucket_name}/data/{base_tbl}/"
        if not s3_location.startswith(expected_s3_prefix) or "'" in s3_location or ";" in s3_location:
            raise PermissionError(f"Invalid S3 table location: {s3_location}")

        validated_cols = []
        for col_name, col_type in col_defs:
            if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_]*", col_name):
                raise ValueError(f"Invalid column identifier: {col_name}")
            upper_type = col_type.strip().upper()
            if not _is_valid_athena_ddl_type(upper_type):
                raise ValueError(f"Unsupported or invalid Athena DDL column type: {col_type}")
            validated_cols.append(f"`{col_name}` {upper_type}")

        cols_sql = ", ".join(validated_cols)
        if storage_format == "PARQUET":
            ddl_sql = (
                f"CREATE EXTERNAL TABLE IF NOT EXISTS {self.database}.{table_name} "
                f"({cols_sql}) "
                f"STORED AS PARQUET "
                f"LOCATION '{s3_location}'"
            )
        elif storage_format == "TEXTFILE":
            if field_delimiter not in (",", "\\t"):
                raise ValueError(f"Unsupported field delimiter: {field_delimiter}")
            ddl_sql = (
                f"CREATE EXTERNAL TABLE IF NOT EXISTS {self.database}.{table_name} "
                f"({cols_sql}) "
                f"ROW FORMAT DELIMITED FIELDS TERMINATED BY '{field_delimiter}' "
                f"STORED AS TEXTFILE "
                f"LOCATION '{s3_location}' "
                f"TBLPROPERTIES ('skip.header.line.count'='1')"
            )
        else:
            raise ValueError(f"Unsupported storage format: {storage_format}")

        parsed = sqlglot.parse(ddl_sql, read="hive", error_level=sqlglot.ErrorLevel.RAISE)
        if len(parsed) != 1 or not isinstance(parsed[0], exp.Create):
            raise PermissionError("AST verification failed for constructed CREATE EXTERNAL TABLE DDL.")

        create_node = parsed[0]
        schema_node = create_node.this
        if not isinstance(schema_node, exp.Schema) or not isinstance(schema_node.this, exp.Table):
            raise PermissionError("CREATE AST is missing expected table schema structure.")

        table_node = schema_node.this
        if (
            table_node.db.lower() != self.database.lower()
            or table_node.name.lower() != table_name.lower()
        ):
            raise PermissionError("CREATE AST target database/table does not match expected identifiers.")

        ast_col_defs = [node for node in schema_node.expressions if isinstance(node, exp.ColumnDef)]
        if len(ast_col_defs) != len(col_defs):
            raise PermissionError(
                f"CREATE AST column count mismatch: expected {len(col_defs)}, got {len(ast_col_defs)}."
            )
        for ast_col, (expected_name, _) in zip(ast_col_defs, col_defs):
            if ast_col.name.lower() != expected_name.lower():
                raise PermissionError(
                    f"CREATE AST column name mismatch: expected '{expected_name}', got '{ast_col.name}'."
                )

        props = create_node.args.get("properties")
        if not isinstance(props, exp.Properties):
            raise PermissionError("CREATE AST is missing table properties.")

        has_external = any(isinstance(p, exp.ExternalProperty) for p in props.expressions)
        loc_literals = [
            lit.this
            for p in props.expressions
            if isinstance(p, exp.LocationProperty)
            for lit in p.find_all(exp.Literal)
        ]
        format_sql_upper = props.sql("hive").upper()

        if not has_external:
            raise PermissionError("CREATE AST is missing EXTERNAL table property.")
        if s3_location not in loc_literals:
            raise PermissionError(f"CREATE AST LOCATION '{loc_literals}' does not match '{s3_location}'.")
        if storage_format not in format_sql_upper:
            raise PermissionError(f"CREATE AST storage format does not match '{storage_format}'.")

        return ddl_sql

    def _execute_ingestion_ddl(
        self,
        operation: str,
        table_name: str,
        col_defs: Optional[List[Tuple[str, str]]] = None,
        storage_format: str = "PARQUET",
        s3_location: str = "",
        field_delimiter: str = ",",
        expected_base_table: Optional[str] = None,
    ) -> str:
        if operation == "drop":
            if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_]*", self.database):
                raise ValueError(f"Invalid database identifier: {self.database}")
            if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_]*", table_name):
                raise ValueError(f"Invalid table identifier: {table_name}")

            ddl_sql = f"DROP TABLE IF EXISTS {self.database}.{table_name}"
            parsed = sqlglot.parse(ddl_sql, read="hive", error_level=sqlglot.ErrorLevel.RAISE)
            if len(parsed) != 1 or not isinstance(parsed[0], exp.Drop):
                raise PermissionError("AST verification failed for constructed DROP TABLE DDL.")
            table_node = parsed[0].this
            if (
                not isinstance(table_node, exp.Table)
                or table_node.db.lower() != self.database.lower()
                or table_node.name.lower() != table_name.lower()
            ):
                raise PermissionError("AST target mismatch in constructed DROP TABLE DDL.")

        elif operation == "create":
            ddl_sql = self._build_and_validate_create_ddl(
                table_name=table_name,
                col_defs=col_defs or [],
                storage_format=storage_format,
                s3_location=s3_location,
                field_delimiter=field_delimiter,
                expected_base_table=expected_base_table,
            )
        else:
            raise PermissionError(f"Unsupported ingestion DDL operation: {operation}")

        try:
            response = self.athena.start_query_execution(
                QueryString=ddl_sql,
                QueryExecutionContext={"Database": self.database},
                ResultConfiguration={"OutputLocation": self.output_location},
                WorkGroup=self.work_group,
            )
        except ClientError as e:
            raise AthenaExecutionError(f"Failed to submit ingestion DDL ({operation}): {e}") from e

        query_execution_id = response["QueryExecutionId"]
        self._poll_athena_execution(query_execution_id)
        return f"Athena Ingestion DDL ({operation.upper()}) Executed: {query_execution_id}"

    def _get_existing_table_snapshot(self, table_name: str) -> Optional[Dict[str, Any]]:
        """
        Captures the current Glue table metadata prior to replacement.
        Returns None ONLY when Glue returns EntityNotFoundException.
        Re-raises AccessDeniedException or any other AWS ClientError as AthenaExecutionError.
        """
        try:
            resp = self.glue.get_table(DatabaseName=self.database, Name=table_name)
            tbl = resp.get("Table", {})
            sd = tbl.get("StorageDescriptor", {})
            location = sd.get("Location", "")
            expected_root = f"s3://{self.bucket_name}/data/{table_name}/"
            s3_prefix = (
                location.replace(f"s3://{self.bucket_name}/", "", 1)
                if isinstance(location, str) and location.startswith(expected_root)
                else None
            )
            cols = [
                (c["Name"], c["Type"].upper())
                for c in sd.get("Columns", [])
                if "Name" in c and "Type" in c
            ]
            input_format = (sd.get("InputFormat") or "").upper()
            storage_format = "PARQUET" if "PARQUET" in input_format else "TEXTFILE"
            serde_params = sd.get("SerdeInfo", {}).get("Parameters", {})
            delim = serde_params.get("field.delim", ",")
            field_delimiter = "\\t" if delim in ("\t", "\\t") else ","

            return {
                "table_name": table_name,
                "s3_location": location,
                "s3_prefix": s3_prefix,
                "col_defs": cols,
                "partition_keys": copy.deepcopy(tbl.get("PartitionKeys", [])),
                "parameters": copy.deepcopy(tbl.get("Parameters", {})),
                "storage_format": storage_format,
                "field_delimiter": field_delimiter,
                "raw_table_definition": copy.deepcopy(tbl),
            }
        except ClientError as e:
            err_code = e.response.get("Error", {}).get("Code", "")
            if err_code == "EntityNotFoundException":
                return None
            raise AthenaExecutionError(
                f"Failed to inspect existing Glue table '{self.database}.{table_name}' "
                f"(AWS Error '{err_code}'): {e}"
            ) from e

    def _verify_table_location_updated(self, table_name: str, expected_s3_location: str) -> None:
        """
        Fail-closed verification via Glue GetTable that the target table's
        StorageDescriptor.Location exactly matches expected_s3_location.
        Used after both forward production cutovers AND rollback restorations.
        """
        try:
            resp = self.glue.get_table(DatabaseName=self.database, Name=table_name)
        except ClientError as e:
            err_code = e.response.get("Error", {}).get("Code", "")
            if err_code == "EntityNotFoundException":
                raise AthenaExecutionError(
                    f"Post-cutover verification failed: Glue table '{self.database}.{table_name}' does not exist."
                ) from e
            raise AthenaExecutionError(
                f"Post-cutover verification failed for '{self.database}.{table_name}': {e}"
            ) from e

        actual_location = (
            resp.get("Table", {})
            .get("StorageDescriptor", {})
            .get("Location")
            if isinstance(resp, dict)
            else None
        )
        if actual_location != expected_s3_location:
            raise AthenaExecutionError(
                f"Cutover verification failed for '{self.database}.{table_name}': "
                f"Glue table metadata was not replaced (expected {expected_s3_location!r}, "
                f"found {actual_location!r})."
            )

    def _delete_s3_prefix(self, s3_prefix: str) -> None:
        """
        Deletes all objects under a failed or retired S3 ingestion prefix and
        inspects the S3 DeleteObjects response for partial deletion errors.
        """
        if not s3_prefix or not s3_prefix.startswith("data/"):
            return
        paginator = self.s3.get_paginator("list_objects_v2")
        for page in paginator.paginate(Bucket=self.bucket_name, Prefix=s3_prefix):
            contents = page.get("Contents", [])
            if contents:
                delete_keys = [{"Key": obj["Key"]} for obj in contents]
                del_resp = self.s3.delete_objects(
                    Bucket=self.bucket_name,
                    Delete={"Objects": delete_keys, "Quiet": False},
                )
                if isinstance(del_resp, dict):
                    partial_errors = del_resp.get("Errors", [])
                    if partial_errors:
                        err_summary = "; ".join(
                            f"{err.get('Key')}: {err.get('Code')} ({err.get('Message')})"
                            for err in partial_errors
                        )
                        logger.error(f"Partial S3 deletion failure under '{s3_prefix}': {err_summary}")
                        raise AthenaExecutionError(
                            f"S3 DeleteObjects reported partial deletion errors under '{s3_prefix}': {err_summary}"
                        )

    def cleanup_abandoned_staging_prefixes(
        self,
        table_name: str,
        min_age_seconds: int = 3600,
        require_verified_state: bool = True,
    ) -> List[str]:
        """
        Deletes only prefixes with verified terminal or lease-expired durable manifests
        that are older than min_age_seconds and are not the active Glue table location.
        Never interprets missing/unverifiable state as abandoned when require_verified_state=True.
        """
        states = self._load_durable_ingestion_states(table_name)
        now = time.time()
        active_location = None
        try:
            snap = self.glue.get_table(DatabaseName=self.database, Name=table_name).get("Table", {})
            active_location = (snap.get("StorageDescriptor") or {}).get("Location")
        except ClientError as e:
            if e.response.get("Error", {}).get("Code") not in ("EntityNotFoundException",):
                raise AthenaExecutionError(f"Cannot verify active Glue location before cleanup: {e}") from e

        active_prefix = (
            str(active_location or "").replace(f"s3://{self.bucket_name}/", "", 1).rstrip("/")
        )
        removed: List[str] = []

        for prefix, rec in sorted(states.items()):
            state = str(rec.get("state", "")).upper()
            age = now - float(rec.get("updated_at", now))
            lease_exp = float(rec.get("lease_expires_at", now + 1))
            lease_expired = lease_exp <= now

            if prefix.rstrip("/") == active_prefix:
                continue
            if age < min_age_seconds:
                continue
            # Retain any prefix whose lease has not expired
            if not lease_expired:
                continue
            # Once lease has expired AND age >= min_age_seconds, allow cleanup of
            # TERMINAL states (FAILED/RETIRED) or crashed in-flight states (UPLOADING/VALIDATING)
            if state not in TERMINAL_INGESTION_STATES and state not in ("UPLOADING", "VALIDATING"):
                continue

            self._delete_s3_prefix(prefix)
            removed.append(prefix)

        return removed

    def get_schema(self) -> str:
        """Retrieves table and column metadata across all Glue Catalog pages for prompting."""
        schema_lines = []
        paginator = self.glue.get_paginator("get_tables")

        for page in paginator.paginate(DatabaseName=self.database):
            for tbl in page.get("TableList", []):
                tbl_name = tbl["Name"]
                if tbl_name.startswith("_staging_"):
                    continue
                cols = [
                    f"{c['Name']} {c['Type']}"
                    for c in tbl.get("StorageDescriptor", {}).get("Columns", [])
                ]
                schema_lines.append(
                    f"CREATE EXTERNAL TABLE {tbl_name} (\n  "
                    + ",\n  ".join(cols)
                    + "\n);"
                )

        return (
            "\n\n".join(schema_lines)
            if schema_lines
            else "-- No tables found in database."
        )

    @staticmethod
    def _sanitize_column_names(raw_columns: List[str]) -> List[str]:
        seen: Dict[str, int] = {}
        clean_cols: List[str] = []
        for col in raw_columns:
            base = re.sub(r"[^a-zA-Z0-9_]", "_", str(col).strip().lower()).strip("_")
            if not base or base[0].isdigit():
                base = f"col_{base}" if base else "col_unnamed"
            if base in seen:
                seen[base] += 1
                unique_col = f"{base}_{seen[base]}"
            else:
                seen[base] = 1
                unique_col = base
            clean_cols.append(unique_col)
        return clean_cols

    @staticmethod
    def _normalize_and_validate_overrides(
        schema_overrides: Optional[Dict[str, str]],
        valid_clean_columns: List[str],
    ) -> Dict[str, str]:
        if not schema_overrides:
            return {}

        valid_col_set = set(valid_clean_columns)
        normalized: Dict[str, str] = {}
        for raw_key, raw_type in schema_overrides.items():
            clean_k = re.sub(r"[^a-zA-Z0-9_]", "_", str(raw_key).strip().lower()).strip("_")
            if not clean_k or clean_k[0].isdigit():
                clean_k = f"col_{clean_k}" if clean_k else "col_unnamed"

            if clean_k in normalized:
                raise ValueError(
                    f"Duplicate schema override column after normalization: '{raw_key}' -> '{clean_k}'."
                )
            if clean_k not in valid_col_set:
                raise ValueError(
                    f"Unknown schema override column '{raw_key}' (normalized: '{clean_k}') "
                    f"not found in dataset columns: {valid_clean_columns}."
                )

            upper_type = str(raw_type).strip().upper()
            if not _is_valid_athena_ddl_type(upper_type):
                raise ValueError(
                    f"Invalid schema override type '{raw_type}' for column '{clean_k}'."
                )
            normalized[clean_k] = upper_type

        return normalized

    @classmethod
    def _infer_and_validate_full_csv_stream(
        cls,
        uploaded_file: Any,
        raw_name: str,
        delimiter: str,
        schema_overrides: Optional[Dict[str, str]] = None,
    ) -> List[Tuple[str, str]]:
        uploaded_file.seek(0)
        text_wrapper = io.TextIOWrapper(uploaded_file, encoding="utf-8", errors="strict", newline="")
        sample_rows: List[List[str]] = []
        try:
            reader = csv.reader(text_wrapper, delimiter=delimiter)
            header = next(reader, None)
            if not header:
                raise ValueError(f"Uploaded CSV/TSV file '{raw_name}' is empty or missing a header row.")

            clean_names = cls._sanitize_column_names(header)
            expected_width = len(clean_names)

            for row_idx, row in enumerate(reader, start=2):
                if len(row) != expected_width:
                    raise ValueError(
                        f"Inconsistent CSV column count at line {row_idx}: header has "
                        f"{expected_width} columns, but row has {len(row)} columns."
                    )
                if len(sample_rows) < 1000:
                    sample_rows.append(row)
        except UnicodeDecodeError as ude:
            raise ValueError(
                f"CSV/TSV file '{raw_name}' contains invalid UTF-8 byte sequences: {ude}"
            ) from ude
        finally:
            try:
                text_wrapper.detach()
            except Exception:
                pass

        return cls._infer_conservative_csv_columns(header, sample_rows, schema_overrides=schema_overrides)

    @classmethod
    def _infer_conservative_csv_columns(
        cls,
        header: List[str],
        rows: List[List[str]],
        schema_overrides: Optional[Dict[str, str]] = None,
    ) -> List[Tuple[str, str]]:
        clean_names = cls._sanitize_column_names(header)
        expected_width = len(clean_names)

        for row_idx, row in enumerate(rows, start=2):
            if len(row) != expected_width:
                raise ValueError(
                    f"Inconsistent CSV column count at line {row_idx}: header has "
                    f"{expected_width} columns, but row has {len(row)} columns."
                )

        caller_overrides = cls._normalize_and_validate_overrides(schema_overrides, clean_names)
        effective_overrides: Dict[str, str] = dict(CUR_COLUMN_TYPE_OVERRIDES)
        effective_overrides.update(caller_overrides)

        int_pattern = re.compile(r"^-?\d+$")
        float_pattern = re.compile(r"^-?(?:\d+\.\d*|\.\d+|\d+(?:[eE][-+]?\d+))$")

        col_defs: List[Tuple[str, str]] = []
        for col_idx, clean_col in enumerate(clean_names):
            if clean_col in effective_overrides:
                col_defs.append((clean_col, effective_overrides[clean_col]))
                continue

            non_empty_vals = [
                row[col_idx].strip()
                for row in rows
                if col_idx < len(row) and row[col_idx] is not None and row[col_idx].strip() != ""
            ]

            if len(non_empty_vals) < 5:
                col_defs.append((clean_col, "STRING"))
                continue

            if any(token in clean_col for token in ("id", "code", "arn", "account", "date", "time", "tag")):
                col_defs.append((clean_col, "STRING"))
                continue

            if all(int_pattern.match(v) for v in non_empty_vals):
                col_defs.append((clean_col, "BIGINT"))
                continue

            if all(int_pattern.match(v) or float_pattern.match(v) for v in non_empty_vals):
                col_defs.append((clean_col, "DOUBLE"))
                continue

            col_defs.append((clean_col, "STRING"))

        return col_defs

    @staticmethod
    def _map_arrow_type_to_athena(arrow_type: pa.DataType, field_name: str = "") -> str:
        if pa.types.is_nested(arrow_type):
            raise ValueError(
                f"Unsupported nested Parquet type '{arrow_type}' on column '{field_name}'. "
                "Please flatten nested structs/lists before ingestion."
            )
        if pa.types.is_uint64(arrow_type):
            raise ValueError(
                f"Unsupported unsigned 64-bit integer (uint64) on column '{field_name}': "
                "Athena BIGINT is signed and may overflow. Cast to DECIMAL(20,0) or INT64 first."
            )
        if pa.types.is_boolean(arrow_type):
            return "BOOLEAN"
        if pa.types.is_int64(arrow_type) or pa.types.is_uint32(arrow_type):
            return "BIGINT"
        if pa.types.is_signed_integer(arrow_type) or pa.types.is_uint8(arrow_type) or pa.types.is_uint16(arrow_type):
            return "INT"
        if pa.types.is_float32(arrow_type):
            return "FLOAT"
        if pa.types.is_float64(arrow_type):
            return "DOUBLE"
        if pa.types.is_decimal(arrow_type):
            dec_type = f"DECIMAL({arrow_type.precision},{arrow_type.scale})"
            if not _validate_athena_decimal_type(dec_type):
                raise ValueError(
                    f"Invalid Parquet DECIMAL({arrow_type.precision},{arrow_type.scale}) on column '{field_name}'."
                )
            return dec_type
        if pa.types.is_date(arrow_type):
            return "DATE"
        if pa.types.is_timestamp(arrow_type):
            return "TIMESTAMP"
        if pa.types.is_string(arrow_type) or pa.types.is_large_string(arrow_type):
            return "STRING"

        raise ValueError(
            f"Unsupported Parquet Arrow data type '{arrow_type}' on column '{field_name}'."
        )

    @classmethod
    def _validate_and_stream_excel_to_temp_csv(
        cls,
        uploaded_file: Any,
        raw_name: str,
        ext: str,
        schema_overrides: Optional[Dict[str, str]] = None,
    ) -> Tuple[str, List[Tuple[str, str]]]:
        uploaded_file.seek(0, io.SEEK_END)
        file_size = uploaded_file.tell()
        uploaded_file.seek(0)
        if file_size > MAX_EXCEL_UPLOAD_BYTES:
            raise ValueError(
                f"Excel file '{raw_name}' ({file_size / (1024 * 1024):.1f} MB) exceeds the "
                f"{MAX_EXCEL_UPLOAD_BYTES // (1024 * 1024)} MB limit for Excel ingestion. "
                "Please convert large workbooks to CSV or Parquet for streaming upload."
            )

        fd, tmp_csv_path = tempfile.mkstemp(suffix=".csv")
        os.close(fd)
        sample_rows: List[List[str]] = []
        header: List[str] = []

        try:
            if ext == ".xlsx":
                wb = openpyxl.load_workbook(uploaded_file, read_only=True, data_only=True)
                try:
                    sheet = wb.active
                    if sheet is None:
                        raise ValueError(f"Excel file '{raw_name}' contains no active worksheet.")
                    if (sheet.max_column or 0) > MAX_EXCEL_COLS or (sheet.max_row or 0) > MAX_EXCEL_ROWS + 1:
                        raise ValueError(
                            f"Excel workbook '{raw_name}' metadata dimensions ({sheet.max_row} rows, "
                            f"{sheet.max_column} columns) exceed safety limits ({MAX_EXCEL_ROWS} rows, {MAX_EXCEL_COLS} columns)."
                        )

                    with open(tmp_csv_path, "w", encoding="utf-8", newline="") as csv_out:
                        writer = csv.writer(csv_out)
                        row_count = 0
                        for row_idx, row_vals in enumerate(sheet.iter_rows(values_only=True), start=1):
                            str_vals = ["" if v is None else str(v) for v in row_vals]
                            if len(str_vals) > MAX_EXCEL_COLS:
                                raise ValueError(
                                    f"Excel row {row_idx} in '{raw_name}' has {len(str_vals)} columns "
                                    f"(limit: {MAX_EXCEL_COLS})."
                                )
                            if row_idx == 1:
                                header = [c for c in str_vals if c != ""] if all(c == "" for c in str_vals) else str_vals
                                if not header:
                                    raise ValueError(f"Uploaded Excel file '{raw_name}' contains no header columns.")
                                writer.writerow(header)
                            else:
                                row_count += 1
                                if row_count > MAX_EXCEL_ROWS:
                                    raise ValueError(
                                        f"Excel workbook '{raw_name}' exceeds row limit of {MAX_EXCEL_ROWS} rows."
                                    )
                                normalized_row = str_vals[: len(header)] + [""] * max(0, len(header) - len(str_vals))
                                writer.writerow(normalized_row)
                                if len(sample_rows) < 1000:
                                    sample_rows.append(normalized_row)
                finally:
                    wb.close()
                    uploaded_file.seek(0)

            elif ext == ".xls":
                raw_bytes = uploaded_file.read()
                uploaded_file.seek(0)
                book = xlrd.open_workbook(file_contents=raw_bytes, on_demand=True)
                try:
                    if book.nsheets == 0:
                        raise ValueError(f"Excel .xls file '{raw_name}' contains no worksheets.")
                    sheet = book.sheet_by_index(0)
                    if sheet.ncols > MAX_EXCEL_COLS or sheet.nrows > MAX_EXCEL_ROWS + 1:
                        raise ValueError(
                            f"Excel .xls workbook '{raw_name}' dimensions ({sheet.nrows} rows, "
                            f"{sheet.ncols} columns) exceed safety limits ({MAX_EXCEL_ROWS} rows, {MAX_EXCEL_COLS} columns)."
                        )
                    if sheet.ncols == 0 or sheet.nrows == 0:
                        raise ValueError(f"Uploaded Excel .xls file '{raw_name}' is empty.")

                    with open(tmp_csv_path, "w", encoding="utf-8", newline="") as csv_out:
                        writer = csv.writer(csv_out)
                        header = [str(sheet.cell_value(0, c)) for c in range(sheet.ncols)]
                        writer.writerow(header)
                        for r_idx in range(1, sheet.nrows):
                            if r_idx > MAX_EXCEL_ROWS:
                                raise ValueError(
                                    f"Excel .xls workbook '{raw_name}' exceeds row limit of {MAX_EXCEL_ROWS} rows."
                                )
                            row_vals = [str(sheet.cell_value(r_idx, c)) for c in range(sheet.ncols)]
                            writer.writerow(row_vals)
                            if len(sample_rows) < 1000:
                                sample_rows.append(row_vals)
                finally:
                    book.release_resources()
            else:
                raise ValueError(f"Unsupported Excel extension: {ext}")

            col_defs = cls._infer_conservative_csv_columns(
                header, sample_rows, schema_overrides=schema_overrides
            )
            return tmp_csv_path, col_defs

        except Exception:
            if os.path.exists(tmp_csv_path):
                os.remove(tmp_csv_path)
            raise

    def _start_lock_heartbeat(self, table_name: str, ingestion_id: str):
        stop_event = threading.Event()
        lost: Dict[str, Optional[Exception]] = {"error": None}
        interval = self.lease_heartbeat_interval_sec

        def beat():
            while not stop_event.wait(interval):
                try:
                    self._renew_distributed_table_lock(table_name, ingestion_id, "HEARTBEAT")
                except Exception as exc:
                    lost["error"] = exc
                    stop_event.set()

        thread = threading.Thread(
            target=beat, name=f"lease-{table_name}-{ingestion_id[:8]}", daemon=True
        )
        thread.start()
        return stop_event, thread, lost

    @staticmethod
    def _stop_lock_heartbeat(heartbeat):
        if heartbeat:
            stop_event, thread, _ = heartbeat
            stop_event.set()
            thread.join(timeout=5)

    @staticmethod
    def _assert_lease_healthy(heartbeat):
        if heartbeat and heartbeat[2].get("error"):
            raise AthenaExecutionError(
                f"Distributed lease heartbeat failed; ingestion aborted: {heartbeat[2]['error']}"
            )

    def stream_upload_and_create_table(
        self,
        uploaded_file: Any,
        schema_overrides: Optional[Dict[str, str]] = None,
        validate_full_dataset: Optional[bool] = None,
    ) -> str:
        """
        Ingests datasets to S3 under a unique versioned prefix (data/<table>/ingestion_id=<uuid>/)
        and performs a fail-closed distributed-locked, heartbeat-renewed, staging-validated
        production cutover.
        """
        raw_name = os.path.basename(getattr(uploaded_file, "name", "") or "")
        if "." not in raw_name:
            raise ValueError(f"Uploaded file '{raw_name}' must have a valid extension (.csv, .tsv, .parquet, .xlsx, .xls).")

        base_stem, raw_ext = raw_name.rsplit(".", 1)
        ext = f".{raw_ext.strip().lower()}"
        tbl_name = re.sub(r"[^a-zA-Z0-9_]", "_", base_stem.strip().lower()).strip("_")
        if not tbl_name or tbl_name[0].isdigit():
            tbl_name = f"tbl_{tbl_name}" if tbl_name else "tbl_dataset"

        safe_s3_filename = f"{tbl_name}{ext}"
        ingestion_id = uuid.uuid4().hex
        staging_tbl_name = f"_staging_{tbl_name}_{ingestion_id[:8]}"
        s3_prefix = f"data/{tbl_name}/ingestion_id={ingestion_id}/"
        s3_location = f"s3://{self.bucket_name}/{s3_prefix}"
        s3_key = f"{s3_prefix}{safe_s3_filename}"
        tmp_excel_csv_path: Optional[str] = None

        table_lock_key = f"{self.database}.{tbl_name}"
        table_lock = self._get_table_lock(table_lock_key)
        if not table_lock.acquire(blocking=False):
            raise AthenaExecutionError(
                f"Concurrent ingestion conflict: an upload is already in progress for table '{table_lock_key}'."
            )

        distributed_lock_acquired = False
        heartbeat = None
        try:
            self._acquire_distributed_table_lock(tbl_name, ingestion_id)
            distributed_lock_acquired = True
            heartbeat = self._start_lock_heartbeat(tbl_name, ingestion_id)

            logger.info(
                f"Starting ingestion '{ingestion_id}' (owner='{self.task_owner_id}') "
                f"for table '{tbl_name}' from file '{raw_name}'."
            )
            previous_snapshot = self._get_existing_table_snapshot(tbl_name)
            previous_s3_prefix = previous_snapshot.get("s3_prefix") if previous_snapshot else None

            # Phase 1: Full structural & schema validation BEFORE uploading to S3
            if ext in (".csv", ".tsv"):
                delimiter = "\t" if ext == ".tsv" else ","
                col_defs = self._infer_and_validate_full_csv_stream(
                    uploaded_file=uploaded_file,
                    raw_name=raw_name,
                    delimiter=delimiter,
                    schema_overrides=schema_overrides,
                )
                storage_format = "TEXTFILE"
                field_delimiter = "\\t" if delimiter == "\t" else ","

            elif ext == ".parquet":
                uploaded_file.seek(0)
                parquet_file = pq.ParquetFile(uploaded_file)
                arrow_schema = parquet_file.schema_arrow
                raw_names = list(arrow_schema.names)
                clean_names = self._sanitize_column_names(raw_names)

                validated_overrides = self._normalize_and_validate_overrides(schema_overrides, clean_names)

                col_defs = []
                for field, clean_col in zip(arrow_schema, clean_names):
                    physical_athena_type = self._map_arrow_type_to_athena(field.type, field_name=field.name)

                    if clean_col in validated_overrides:
                        requested_type = validated_overrides[clean_col]
                        allowed_targets = COMPATIBLE_PARQUET_TYPE_OVERRIDES.get(
                            physical_athena_type, {physical_athena_type}
                        )
                        if requested_type not in allowed_targets:
                            raise ValueError(
                                f"Incompatible Parquet schema override for column '{clean_col}': "
                                f"cannot override physical Parquet type '{physical_athena_type}' ({field.type}) "
                                f"to '{requested_type}'."
                            )
                        athena_type = requested_type
                    elif clean_col in CUR_COLUMN_TYPE_OVERRIDES and pa.types.is_floating(field.type):
                        athena_type = CUR_COLUMN_TYPE_OVERRIDES[clean_col]
                    else:
                        athena_type = physical_athena_type

                    col_defs.append((clean_col, athena_type))

                storage_format = "PARQUET"
                field_delimiter = ","

            elif ext in (".xlsx", ".xls"):
                tmp_excel_csv_path, col_defs = self._validate_and_stream_excel_to_temp_csv(
                    uploaded_file=uploaded_file,
                    raw_name=raw_name,
                    ext=ext,
                    schema_overrides=schema_overrides,
                )
                storage_format = "TEXTFILE"
                field_delimiter = ","
                s3_key = f"{s3_prefix}{tbl_name}.csv"
            else:
                raise ValueError(f"Unsupported file format: {raw_name}")

            self._build_and_validate_create_ddl(
                table_name=staging_tbl_name,
                col_defs=col_defs,
                storage_format=storage_format,
                s3_location=s3_location,
                field_delimiter=field_delimiter,
                expected_base_table=tbl_name,
            )
            self._build_and_validate_create_ddl(
                table_name=tbl_name,
                col_defs=col_defs,
                storage_format=storage_format,
                s3_location=s3_location,
                field_delimiter=field_delimiter,
            )

            # Phase 2: Upload to isolated S3 prefix, register & validate staging table, then cut over
            table_modified_or_dropped = False
            staging_created = False
            self._set_ingestion_state(s3_prefix, tbl_name, ingestion_id, "UPLOADING")

            try:
                if ext in (".csv", ".tsv"):
                    uploaded_file.seek(0)
                    self.s3.upload_fileobj(
                        uploaded_file,
                        self.bucket_name,
                        s3_key,
                        Config=self.transfer_config,
                    )
                elif ext == ".parquet":
                    if raw_names == clean_names:
                        uploaded_file.seek(0)
                        self.s3.upload_fileobj(
                            uploaded_file,
                            self.bucket_name,
                            s3_key,
                            Config=self.transfer_config,
                        )
                    else:
                        renamed_fields = [
                            field.with_name(clean_col)
                            for field, clean_col in zip(arrow_schema, clean_names)
                        ]
                        renamed_schema = pa.schema(renamed_fields)
                        fd, tmp_path = tempfile.mkstemp(suffix=".parquet")
                        os.close(fd)
                        try:
                            with pq.ParquetWriter(tmp_path, renamed_schema) as writer:
                                for batch in parquet_file.iter_batches(batch_size=50_000):
                                    renamed_batch = pa.RecordBatch.from_arrays(
                                        batch.columns, schema=renamed_schema
                                    )
                                    writer.write_batch(renamed_batch)
                            with open(tmp_path, "rb") as tmp_file_obj:
                                self.s3.upload_fileobj(
                                    tmp_file_obj,
                                    self.bucket_name,
                                    s3_key,
                                    Config=self.transfer_config,
                                )
                        finally:
                            if os.path.exists(tmp_path):
                                os.remove(tmp_path)
                elif ext in (".xlsx", ".xls") and tmp_excel_csv_path:
                    with open(tmp_excel_csv_path, "rb") as tmp_csv_obj:
                        self.s3.upload_fileobj(
                            tmp_csv_obj,
                            self.bucket_name,
                            s3_key,
                            Config=self.transfer_config,
                        )

                # Step 2A: Register temporary staging Glue table & run bounded multi-column validation
                self._assert_lease_healthy(heartbeat)
                self._set_ingestion_state(s3_prefix, tbl_name, ingestion_id, "VALIDATING")
                self._execute_ingestion_ddl(
                    operation="create",
                    table_name=staging_tbl_name,
                    col_defs=col_defs,
                    storage_format=storage_format,
                    s3_location=s3_location,
                    field_delimiter=field_delimiter,
                    expected_base_table=tbl_name,
                )
                staging_created = True

                self._validate_staging_table(
                    staging_tbl_name, col_defs, validate_full_dataset=validate_full_dataset
                )

                # Step 2B: Production cutover (fenced by lease check)
                self._assert_lease_healthy(heartbeat)
                self._set_ingestion_state(s3_prefix, tbl_name, ingestion_id, "CUTOVER")
                if previous_snapshot is not None and self.use_atomic_glue_cutover:
                    table_modified_or_dropped = True
                    self._atomic_glue_update_table(
                        table_name=tbl_name,
                        col_defs=col_defs,
                        storage_format=storage_format,
                        s3_location=s3_location,
                        field_delimiter=field_delimiter,
                        existing_snapshot=previous_snapshot,
                    )
                else:
                    if previous_snapshot is not None:
                        self._execute_ingestion_ddl(operation="drop", table_name=tbl_name)
                        table_modified_or_dropped = True

                    self._execute_ingestion_ddl(
                        operation="create",
                        table_name=tbl_name,
                        col_defs=col_defs,
                        storage_format=storage_format,
                        s3_location=s3_location,
                        field_delimiter=field_delimiter,
                    )

                # Fail-closed verification and ACTIVE transition run for BOTH atomic and non-atomic paths
                self._assert_lease_healthy(heartbeat)
                self._verify_table_location_updated(tbl_name, s3_location)
                self._set_ingestion_state(s3_prefix, tbl_name, ingestion_id, "ACTIVE")

            except Exception as ingestion_err:
                try:
                    self._set_ingestion_state(s3_prefix, tbl_name, ingestion_id, "FAILED")
                except Exception:
                    pass
                logger.error(
                    f"Ingestion '{ingestion_id}' failed for table '{tbl_name}': {ingestion_err}. "
                    "Initiating staging cleanup & verified rollback."
                )

                if staging_created:
                    try:
                        self._execute_ingestion_ddl(operation="drop", table_name=staging_tbl_name)
                        staging_created = False
                    except Exception as stg_drop_err:
                        logger.warning(f"Failed to drop temporary staging table '{staging_tbl_name}': {stg_drop_err}")

                try:
                    self._delete_s3_prefix(s3_prefix)
                except Exception as s3_cleanup_err:
                    logger.error(f"Failed to clean up staging S3 prefix '{s3_prefix}': {s3_cleanup_err}")

                # Restore previous production table AND independently verify restored Glue Location
                if (
                    table_modified_or_dropped
                    and previous_snapshot
                    and previous_snapshot.get("col_defs")
                    and previous_snapshot.get("s3_location")
                ):
                    expected_prev_loc = previous_snapshot["s3_location"]
                    try:
                        if self.use_atomic_glue_cutover:
                            self._atomic_glue_update_table(
                                table_name=tbl_name,
                                col_defs=previous_snapshot["col_defs"],
                                storage_format=previous_snapshot["storage_format"],
                                s3_location=expected_prev_loc,
                                field_delimiter=previous_snapshot["field_delimiter"],
                                existing_snapshot=previous_snapshot,
                            )
                        else:
                            self._execute_ingestion_ddl(
                                operation="create",
                                table_name=tbl_name,
                                col_defs=previous_snapshot["col_defs"],
                                storage_format=previous_snapshot["storage_format"],
                                s3_location=expected_prev_loc,
                                field_delimiter=previous_snapshot["field_delimiter"],
                            )

                        # Independent fail-closed verification of restored production location
                        self._verify_table_location_updated(tbl_name, expected_prev_loc)

                    except Exception as restore_err:
                        crit_msg = (
                            f"CRITICAL RECOVERY FAILURE for table '{self.database}.{tbl_name}' "
                            f"(ingestion_id='{ingestion_id}'): replacement failed ({ingestion_err}) "
                            f"AND rollback/verification to previous snapshot '{expected_prev_loc}' "
                            f"failed ({restore_err})."
                        )
                        logger.critical(crit_msg)
                        raise AthenaExecutionError(crit_msg) from restore_err

                raise ingestion_err

            finally:
                if staging_created:
                    try:
                        self._execute_ingestion_ddl(operation="drop", table_name=staging_tbl_name)
                    except Exception as stg_cleanup_err:
                        logger.warning(
                            f"Non-fatal: could not drop temporary staging table '{staging_tbl_name}': {stg_cleanup_err}"
                        )

            # Phase 3: Clean up retired S3 prefix ONLY after new table cutover & location verification succeed
            if previous_s3_prefix and previous_s3_prefix != s3_prefix:
                try:
                    self._set_ingestion_state(previous_s3_prefix, tbl_name, ingestion_id, "RETIRED")
                    self._delete_s3_prefix(previous_s3_prefix)
                except Exception as cleanup_err:
                    logger.warning(f"Non-fatal: could not remove retired S3 prefix '{previous_s3_prefix}': {cleanup_err}")

            logger.info(f"Completed ingestion '{ingestion_id}' for table '{tbl_name}' at '{s3_location}'.")
            return (
                f"Dataset `{raw_name}` successfully ingested to `{s3_location}` and registered as table"
                f" `{tbl_name}` in Glue Catalog!"
            )

        finally:
            if tmp_excel_csv_path and os.path.exists(tmp_excel_csv_path):
                try:
                    os.remove(tmp_excel_csv_path)
                except Exception:
                    pass
            self._stop_lock_heartbeat(heartbeat)
            if distributed_lock_acquired:
                self._release_distributed_table_lock(tbl_name, ingestion_id)
            table_lock.release()
